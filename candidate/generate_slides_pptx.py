import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    BG_DARK = RGBColor(15, 22, 41)       # #0f1629
    CARD_BG = RGBColor(26, 31, 53)       # #1a1f35
    ACCENT_PURPLE = RGBColor(99, 102, 241) # #6366f1
    TEXT_LIGHT = RGBColor(241, 245, 249) # #f1f5f9
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8
    BORDER_COLOR = RGBColor(40, 48, 77)
    GREEN = RGBColor(16, 185, 129)
    AMBER = RGBColor(245, 158, 11)

    slides_data = [
        {
            "num": "01",
            "title": "The Manager Problem & Key Findings",
            "subtitle": "Navigating noisy calls, late transcripts, and conflicting CRM signals",
            "cards": [
                {
                    "title": "The Manager Dilemma",
                    "bullets": [
                        "Revenue leaders receive fragmented calls, calendar invites, and unstandardized notes across siloed tools.",
                        "Transcripts arrive asynchronously or fail to record entirely.",
                        "Stored CRM labels frequently contradict prospect statements in real conversations."
                    ]
                },
                {
                    "title": "Key Findings (21 Aug 2026, Kolkata)",
                    "bullets": [
                        "44 logical interactions: 30 calls + 14 meetings.",
                        "Text coverage: Only 18 records (40.9%) have usable text; 26 records lack text and must not have fake summaries invented.",
                        "Call connection rate: 36.7% (11 of 30 calls connected)."
                    ]
                },
                {
                    "title": "Top Three Actionable Signals",
                    "bullets": [
                        "M07 (Willowfield): 4 severe software bugs reported (turns 2, 8, 16, 22); escalated to dev team.",
                        "M03 (Asterloom): Pricing quote conflict (INR 26k vs 39k document) surfaced in Turn 24.",
                        "C10 (Meadowbrick): Contact departed 6 months ago (Turn 1); CRM list updated."
                    ]
                }
            ]
        },
        {
            "num": "02",
            "title": "Interface Choices & Deliberate Omissions",
            "subtitle": "Designing for cognitive clarity, rapid triage, and high trust",
            "cards": [
                {
                    "title": "High-Density Command Center",
                    "bullets": [
                        "Dark-mode SaaS workspace with glowing status badges and responsive layout.",
                        "Top-level metric cards act as live filters (e.g. clicking 'Needs Review' isolates exceptions).",
                        "Clean visual separation between calls and meetings, avoiding misleading aggregate conflation."
                    ]
                },
                {
                    "title": "Slide-Out Conversation Inspector",
                    "bullets": [
                        "Context-preserving drawer displays full turn-by-turn dialogue alongside extracted insights.",
                        "Explicit speaker attribution with timecodes for meetings.",
                        "Dedicated coaching notes and next-step ownership blocks."
                    ]
                },
                {
                    "title": "What We Deliberately Left Out",
                    "bullets": [
                        "Speculative pipeline dollar values: no CRM deal table was provided, so guessing amounts produces false precision.",
                        "Cross-rep rankings: comparing outbound cold callers to demo meeting organizers is statistically invalid.",
                        "Distracting decorative charts that do not inform immediate managerial action."
                    ]
                }
            ]
        },
        {
            "num": "03",
            "title": "Routing Architecture & The Difficult Case",
            "subtitle": "Evidence-backed taxonomy preventing pipeline contamination",
            "cards": [
                {
                    "title": "Routing Taxonomy",
                    "bullets": [
                        "Sales Prospecting (cold outreach qualification)",
                        "Sales Follow-up (active deal progression)",
                        "Partnership Development (referral networks)",
                        "Customer Success (support, bugs, onboarding)",
                        "Vendor Evaluation (SpendNest procuring)",
                        "Product Discovery (in-depth prospect demo)",
                        "Internal Training (team knowledge transfer)",
                        "Needs Review (missing or conflicting data)"
                    ]
                },
                {
                    "title": "The Difficult Case: M02 (LedgerBridge)",
                    "bullets": [
                        "Ambiguity: Initial calendar metadata resembled an outbound sales meeting.",
                        "Evidence: Transcript inspection showed SpendNest internal leadership negotiating issuing rails and pricing from LedgerBridge.",
                        "Resolution: Routed to Vendor Evaluation with procurement next steps, keeping sales quota metrics unpolluted."
                    ]
                },
                {
                    "title": "Export Organization",
                    "bullets": [
                        "Structured output directory: exports/manifest.json/csv, exports/analysis/*.json, exports/manager_brief.md.",
                        "Every manifest path is verified to exist on disk.",
                        "1-click ZIP archive generator built into the UI."
                    ]
                }
            ]
        },
        {
            "num": "04",
            "title": "Verification of Numbers & Evidence Tracing",
            "subtitle": "Mathematically sound denominators and quote-level auditability",
            "cards": [
                {
                    "title": "Rigorous Numerical Checks",
                    "bullets": [
                        "Composite Identity: Keyed strictly by (source_type, source_id).",
                        "Idempotency Proof: Re-importing batch 1 processes 44 records: 0 new, 0 updated, 44 skipped.",
                        "Revision Ordering: Batch 2 updates M07 from rev 1 to 2 without adding a 45th record.",
                        "Rep Distinction: 5 call reps separated from meeting organizers."
                    ]
                },
                {
                    "title": "Auditable Evidence Anchors",
                    "bullets": [
                        "M04 Churn threat: 'If we cannot resolve the credit limit issue, we will have to evaluate alternative platforms.' (Turn 19)",
                        "M09 Login failure: 'The token link sent by the system had expired by the time our IT team attempted to authenticate.' (Turn 4)",
                        "C01 Rejection: 'I didn\\'t ask for this call, and I\\'m not interested.' (Quote)"
                    ]
                },
                {
                    "title": "Handling Missing Material",
                    "bullets": [
                        "26 records lack usable text (e.g. B01-B20, cancelled notetakers).",
                        "System displays explicit intake reasons rather than generating hallucinated summaries.",
                        "All 44 records maintain clear review and freshness states."
                    ]
                }
            ]
        },
        {
            "num": "05",
            "title": "Trade-offs, AI Usage, Limitations & Roadmap",
            "subtitle": "Practical engineering decisions and paths to production",
            "cards": [
                {
                    "title": "AI Usage & Determinism",
                    "bullets": [
                        "Pre-generated analysis packaged in analysis.js for 100% offline reliability, speed, and privacy.",
                        "Zero runtime token costs, credentials, or API failure points.",
                        "Late transcript updates trigger stale flags for human managerial sign-off."
                    ]
                },
                {
                    "title": "Architectural Trade-offs",
                    "bullets": [
                        "Vanilla JS & LocalStorage: chosen over complex Docker/Node stacks to run immediately on any reviewer machine.",
                        "Manual route overrides survive re-imports, browser refreshes, and late updates.",
                        "Limitation: Single-user browser storage; production would use PostgreSQL."
                    ]
                },
                {
                    "title": "Production Roadmap",
                    "bullets": [
                        "Multi-tenant Cloud DB with row-level security and audit logging.",
                        "Real-time audio streaming with Whisper/Deepgram STT.",
                        "Bi-directional CRM sync (Salesforce/HubSpot) to automatically assign follow-ups and update stages."
                    ]
                }
            ]
        }
    ]

    for s_idx, data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)

        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()

        # Header Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(1.0), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT_PURPLE
        badge.line.fill.background()
        tf_badge = badge.text_frame
        tf_badge.text = f"SLIDE {data['num']}"
        for p in tf_badge.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT
            p.alignment = PP_ALIGN.CENTER

        # Title
        txBox = slide.shapes.add_textbox(Inches(2.0), Inches(0.4), Inches(10.5), Inches(0.6))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data["title"]
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

        # Subtitle
        txBox_sub = slide.shapes.add_textbox(Inches(2.0), Inches(0.95), Inches(10.5), Inches(0.4))
        tf_sub = txBox_sub.text_frame
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = TEXT_MUTED

        # 3 Columns Cards
        card_width = Inches(3.64)
        card_height = Inches(5.2)
        card_gap = Inches(0.4)
        left_start = Inches(0.8)
        top_start = Inches(1.55)

        for c_idx, card in enumerate(data["cards"]):
            c_left = left_start + c_idx * (card_width + card_gap)
            c_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, top_start, card_width, card_height)
            c_box.fill.solid()
            c_box.fill.fore_color.rgb = CARD_BG
            c_box.line.color.rgb = BORDER_COLOR
            c_box.line.width = Pt(1.5)

            # Card Header
            header_box = slide.shapes.add_textbox(c_left + Inches(0.2), top_start + Inches(0.15), card_width - Inches(0.4), Inches(0.7))
            h_tf = header_box.text_frame
            h_tf.word_wrap = True
            h_p = h_tf.paragraphs[0]
            h_p.text = card["title"]
            h_p.font.size = Pt(15)
            h_p.font.bold = True
            h_p.font.color.rgb = RGBColor(129, 140, 248)

            # Card Bullets
            body_box = slide.shapes.add_textbox(c_left + Inches(0.2), top_start + Inches(0.85), card_width - Inches(0.4), card_height - Inches(1.0))
            b_tf = body_box.text_frame
            b_tf.word_wrap = True
            
            for b_idx, bullet in enumerate(card["bullets"]):
                bp = b_tf.add_paragraph() if b_idx > 0 else b_tf.paragraphs[0]
                bp.text = "• " + bullet
                bp.font.size = Pt(12)
                bp.font.color.rgb = TEXT_LIGHT
                bp.space_after = Pt(12)

    prs.save("slides.pptx")
    print("Successfully generated slides.pptx (5 widescreen slides)")

if __name__ == "__main__":
    create_deck()
